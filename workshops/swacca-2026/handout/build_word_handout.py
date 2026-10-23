"""Create an editable Word companion using the packet's approved content.

Requires python-docx, Pillow and ReportLab. The PDF remains the visual reference.
Word tables, paragraphs, styles, page sections and links remain editable.
"""
from pathlib import Path
from io import BytesIO
import json
import argparse

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw
from reportlab.graphics.barcode.qr import QrCodeWidget

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
WIDTH = 510


def elem(parent, tag, attrs=None):
    e = OxmlElement('w:' + tag)
    for key, value in (attrs or {}).items():
        e.set(qn('w:' + key), str(value))
    parent.append(e)
    return e


def style(doc, name, size, leading, after=0, before=0, bold=False, font='Segoe UI'):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, 1)
    s.font.name = font
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.underline = False
    s.font.color.rgb = RGBColor.from_string('151515')
    rpr = s.element.get_or_add_rPr()
    for node in list(rpr):
        if node.tag in [qn('w:spacing'), qn('w:kern'), qn('w:iCs'), qn('w:szCs')]:
            rpr.remove(node)
    ppr = s.element.get_or_add_pPr()
    for node in list(ppr):
        if node.tag in [qn('w:pBdr'), qn('w:numPr'), qn('w:contextualSpacing')]:
            ppr.remove(node)
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = elem(rpr, 'rFonts')
    for k in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + k), font)
    for k in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + k), None)
    p = s.paragraph_format
    p.line_spacing = Pt(leading)
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.widow_control = False
    return s


def para(parent, text='', sty='Normal'):
    return parent.add_paragraph(text, style=sty)


def spacer(doc, height):
    p = para(doc)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(height)
    p.add_run().font.size = Pt(1)
    return p


def border(parent, tag='pBdr', edges=('bottom',), color='AAAAAA', size=4, space=0):
    node = elem(parent, tag)
    for side in edges:
        elem(node, side, {'val': 'single', 'color': color, 'sz': size, 'space': space})
    return node


def setup_table(parent, widths, padding=(6, 7, 6, 7)):
    t = parent.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    pr = t._tbl.tblPr
    for node in list(pr):
        if node.tag in [qn('w:tblW'), qn('w:tblBorders'), qn('w:tblCellMar')]:
            pr.remove(node)
    elem(pr, 'tblW', {'type': 'dxa', 'w': round(sum(widths) * 20)})
    elem(pr, 'tblInd', {'type': 'dxa', 'w': round(padding[3] * 20)})
    mar = elem(pr, 'tblCellMar')
    for side, value in zip(['top', 'left', 'bottom', 'right'], padding):
        elem(mar, side, {'type': 'dxa', 'w': round(value * 20)})
    borders = elem(pr, 'tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        elem(borders, edge, {'val': 'nil'})
    for c, width in zip(t.columns, widths):
        c.width = Pt(width)
    return t


def row(t, texts, style_name='Table text', shade=False, bottom=None):
    r = t.add_row()
    elem(r._tr.get_or_add_trPr(), 'cantSplit')
    for c, text in zip(r.cells, texts):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.style = style_name
        p.add_run(text)
        pr = c._tc.get_or_add_tcPr()
        if shade:
            elem(pr, 'shd', {'fill': 'F3F3F3'})
        if bottom:
            border(pr, 'tcBorders', color=bottom, size=4)
    return r


def link(p, label, url, bold=False):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = elem(h, 'r')
    pr = elem(r, 'rPr')
    elem(pr, 'color', {'val': '151515'})
    elem(pr, 'u', {'val': 'none'})
    if bold:
        elem(pr, 'b')
    elem(r, 't').text = label
    p._p.append(h)


def field(p, code, cached='1'):
    r = p.add_run()
    elem(r._r, 'fldChar', {'fldCharType': 'begin'})
    ins = elem(r._r, 'instrText')
    ins.set(qn('xml:space'), 'preserve')
    ins.text = ' ' + code + ' '
    elem(r._r, 'fldChar', {'fldCharType': 'separate'})
    elem(r._r, 't').text = cached
    elem(r._r, 'fldChar', {'fldCharType': 'end'})


def qr_image(url):
    qr = QrCodeWidget(url, barLevel='M')
    qr.qr.make()
    grid = qr.qr.modules
    pixels = (len(grid) + 8) * 8
    img = Image.new('RGB', (pixels, pixels), 'white')
    draw = ImageDraw.Draw(img)
    for y, line in enumerate(grid):
        for x, dark in enumerate(line):
            if dark:
                draw.rectangle(((x+4)*8, (y+4)*8, (x+5)*8-1, (y+5)*8-1), fill='black')
    stream = BytesIO()
    img.save(stream, format='PNG')
    stream.seek(0)
    return stream


def section_header(section, label, edition, number):
    section.page_width = Pt(612)
    section.page_height = Pt(792)
    section.left_margin = section.right_margin = Pt(51)
    section.top_margin = Pt(68)
    section.bottom_margin = Pt(56)
    section.header_distance = Pt(27)
    section.footer_distance = Pt(25)
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.style = 'Running header'
    hp.paragraph_format.left_indent = Pt(-6)
    hp.paragraph_format.right_indent = Pt(-6)
    hp.paragraph_format.tab_stops.add_tab_stop(Pt(516), WD_TAB_ALIGNMENT.RIGHT)
    hp.add_run('BIM AGENT LAB / SWACCA 2026\t' + label)
    border(hp._p.get_or_add_pPr(), color='151515', size=5, space=7)
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    fp.style = 'Running footer'
    fp.paragraph_format.left_indent = Pt(-6)
    fp.paragraph_format.right_indent = Pt(-6)
    fp.paragraph_format.tab_stops.add_tab_stop(Pt(516), WD_TAB_ALIGNMENT.RIGHT)
    border(fp._p.get_or_add_pPr(), edges=('top',), space=7)
    fp.add_run(edition + '\t')
    field(fp, 'PAGE', str(number))
    fp.add_run(' / ')
    field(fp, 'NUMPAGES', '8')


def build(output):
    data = json.loads((HERE / 'content.json').read_text(encoding='utf-8'))
    doc = Document()
    doc.core_properties.title = data['title']
    doc.core_properties.subject = 'SWACCA BIM Summit October 29-30 2026 participant handout'
    doc.core_properties.author = 'BIM Agent Lab'
    doc.core_properties.last_modified_by = 'BIM Agent Lab'
    doc.core_properties.comments = ''
    style(doc, 'Normal', 10.5, 14.3, after=9)
    style(doc, 'Title', 25, 29, after=7, bold=True)
    style(doc, 'Subtitle', 11, 15, after=16)
    style(doc, 'Heading 1', 12.5, 16, before=4, after=6, bold=True)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    style(doc, 'Small', 8.5, 11.5, after=6)
    style(doc, 'Table text', 9.5, 12.5)
    style(doc, 'Table heading', 9.5, 12, bold=True)
    style(doc, 'Prompt heading', 9.5, 12, after=5, bold=True)
    style(doc, 'Code', 10, 15, after=9, font='Consolas')
    style(doc, 'Running header', 8, 10, bold=True)
    style(doc, 'Running footer', 8, 10)
    for name in ['Subtitle', 'Small', 'Running header', 'Running footer']:
        doc.styles[name].font.color.rgb = RGBColor.from_string('555555')
    for s in ['Title', 'Subtitle']:
        doc.styles[s].paragraph_format.keep_with_next = True
    # Prevent Word's document-grid snapping from altering chosen line spacing.
    elem(doc.styles['Normal'].element.get_or_add_pPr(), 'snapToGrid', {'val': '0'})
    settings = doc.settings.element
    elem(settings, 'updateFields', {'val': 'true'})

    for number, page in enumerate(data['pages'], 1):
        sec = doc.sections[0] if number == 1 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        if number > 1:
            end = doc.paragraphs[-1]
            end.paragraph_format.line_spacing = Pt(1)
            end.paragraph_format.space_after = Pt(0)
            end.paragraph_format.space_before = Pt(0)
        section_header(sec, page['label'], data['edition'], number)
        para(doc, page['title'], 'Title')
        para(doc, page['subtitle'], 'Subtitle')
        for b in page['blocks']:
            kind = b['type']
            if kind in ['p', 'h2', 'small', 'code']:
                para(doc, b['text'], {'p':'Normal','h2':'Heading 1','small':'Small','code':'Code'}[kind])
            elif kind in ['callout', 'prompt']:
                t = setup_table(doc, [WIDTH], padding=(9, 11, 8, 11))
                for existing in list(t._tbl.tblPr):
                    if existing.tag == qn('w:tblBorders'):
                        t._tbl.tblPr.remove(existing)
                border(t._tbl.tblPr, 'tblBorders', edges=('top','left','bottom','right'), color='151515', size=5)
                r = row(t, [''], 'Prompt heading', shade=kind=='callout')
                c = r.cells[0]
                c.paragraphs[0].add_run(b['title'])
                p = para(c, b['text'])
                p.paragraph_format.space_after = Pt(0)
                border(c._tc.get_or_add_tcPr(), 'tcBorders', edges=('top','left','bottom','right'), color='151515', size=5)
                spacer(doc, 11)
            elif kind == 'table':
                t = setup_table(doc, [WIDTH*x for x in b['widths']])
                head = row(t, b['headers'], 'Table heading', shade=True, bottom='151515')
                elem(head._tr.get_or_add_trPr(), 'tblHeader')
                for values in b['rows']:
                    row(t, values, bottom='BBBBBB')
                spacer(doc, 11)
            elif kind == 'steps':
                for i, text in enumerate(b['items'], 1):
                    para(doc, f'{i}. {text}')
            elif kind == 'lines':
                p = para(doc, b['title'], 'Table heading')
                p.paragraph_format.keep_with_next = True
                for _ in range(b['count']):
                    p = para(doc)
                    p.paragraph_format.line_spacing = Pt(21)
                    p.paragraph_format.space_after = Pt(0)
                    border(p._p.get_or_add_pPr(), edges=('bottom', 'between'), size=3)
                spacer(doc, 8)
            elif kind == 'flow':
                t = setup_table(doc, [152,27,152,27,152], padding=(10,6,10,6))
                r = row(t, [b['items'][0],'>',b['items'][1],'>',b['items'][2]], 'Table heading')
                for i in [0,2,4]:
                    border(r.cells[i]._tc.get_or_add_tcPr(), 'tcBorders', edges=('top','left','bottom','right'), color='151515')
                spacer(doc, 10)
            elif kind == 'repo':
                t = setup_table(doc, [408,102], padding=(0,0,0,0))
                r = row(t, ['', ''])
                p = r.cells[0].paragraphs[0]
                p.style = 'Normal'
                link(p, b['text'], data['repo_url'], bold=True)
                para(r.cells[0], b['caption'], 'Small')
                p = r.cells[1].paragraphs[0]
                p.paragraph_format.line_spacing = 1
                pic = p.add_run().add_picture(qr_image(data['repo_url']), width=Pt(76))
                pic._inline.docPr.set('descr', 'QR code for ' + data['repo_url'])
                spacer(doc, 9)
            elif kind == 'references':
                for item in b['items']:
                    p = para(doc, sty='Small')
                    link(p, item['label'], item['url'])
            else:
                raise ValueError(kind)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f'Created editable Word handout: {output}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=ROOT/'output/docx/swacca-2026-participant-handout-oct29-30.docx')
    build(parser.parse_args().output)
